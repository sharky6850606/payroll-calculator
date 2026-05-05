from __future__ import annotations

import copy
import re
from datetime import datetime
from io import BytesIO
from typing import Iterable

import openpyxl
import streamlit as st
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.datavalidation import DataValidationList
from openpyxl.styles import PatternFill


TARGET_SHEET = "Original record"
OUTPUT_FILENAME = "payroll_original_record_updated.xlsx"
TIME_RE = re.compile(r"(?<!\d)([01]?\d|2[0-3]):([0-5]\d)(?!\d)")
GREEN_FILL = PatternFill(fill_type="solid", fgColor="70AD47")


def extract_time_minutes(value: object) -> list[int]:
    """Return valid HH:MM times from a cell as minutes after midnight."""
    if value is None:
        return []

    text = str(value)
    minutes: list[int] = []
    for match in TIME_RE.finditer(text):
        hour = int(match.group(1))
        minute = int(match.group(2))
        minutes.append(hour * 60 + minute)
    return minutes


def calculate_hours_from_times(times: list[int]) -> float | None:
    """Treat times as in/out pairs and return decimal hours."""
    if len(times) < 2:
        return None

    total_minutes = 0
    for index in range(0, len(times) - 1, 2):
        clock_in = times[index]
        clock_out = times[index + 1]
        if clock_out < clock_in:
            clock_out += 24 * 60
        total_minutes += clock_out - clock_in

    return total_minutes / 60


def is_employee_header_row(ws: openpyxl.worksheet.worksheet.Worksheet, row: int) -> bool:
    for cell in ws[row]:
        if isinstance(cell.value, str) and "Name:" in cell.value:
            return True
    return False


def find_employee_header_rows(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[int]:
    return [row for row in range(1, ws.max_row + 1) if is_employee_header_row(ws, row)]


def find_next_employee_header(ws: openpyxl.worksheet.worksheet.Worksheet, start_row: int) -> int | None:
    for row in range(start_row + 1, ws.max_row + 1):
        if is_employee_header_row(ws, row):
            return row
    return None


def is_date_header(value: object) -> bool:
    if value is None:
        return False

    if isinstance(value, datetime):
        return True

    if isinstance(value, (int, float)) and float(value).is_integer():
        return 1 <= int(value) <= 31

    text = str(value).strip()
    if re.fullmatch(r"\d{1,2}", text):
        return 1 <= int(text) <= 31

    return bool(re.fullmatch(r"\d{1,2}[/-]\d{1,2}([/-]\d{2,4})?", text))


def find_date_columns(ws: openpyxl.worksheet.worksheet.Worksheet, date_row: int) -> list[int]:
    return [
        cell.column
        for cell in ws[date_row]
        if not isinstance(cell, MergedCell) and is_date_header(cell.value)
    ]


def color_to_rgb(cell: openpyxl.cell.cell.Cell) -> str:
    color = cell.fill.fgColor
    if color.type == "rgb" and color.rgb:
        return color.rgb[-6:].upper()
    if color.type == "indexed" and color.indexed is not None:
        indexed = openpyxl.styles.colors.COLOR_INDEX[color.indexed]
        return indexed[-6:].upper()
    return ""


def is_green_cell(cell: openpyxl.cell.cell.Cell) -> bool:
    rgb = color_to_rgb(cell)
    if len(rgb) != 6:
        return False

    red = int(rgb[0:2], 16)
    green = int(rgb[2:4], 16)
    blue = int(rgb[4:6], 16)
    return green >= 110 and green > red * 1.15 and green > blue * 1.15


def row_has_green_summary_fill(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, columns: Iterable[int]
) -> bool:
    checked = [ws.cell(row=row, column=column) for column in columns]
    return any(is_green_cell(cell) for cell in checked)


def find_green_summary_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    end_row: int,
    date_columns: list[int],
) -> int | None:
    for row in range(start_row, end_row + 1):
        if row_has_green_summary_fill(ws, row, date_columns):
            return row
    return None


def row_has_attendance_times(
    ws: openpyxl.worksheet.worksheet.Worksheet, row: int, date_columns: list[int]
) -> bool:
    return any(extract_time_minutes(ws.cell(row=row, column=column).value) for column in date_columns)


def find_first_blank_attendance_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    end_row: int,
    date_columns: list[int],
) -> int | None:
    for row in range(start_row, end_row + 1):
        if not row_has_attendance_times(ws, row, date_columns):
            return row
    return None


def copy_row_style(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    source_row: int,
    target_row: int,
    max_column: int,
) -> None:
    for column in range(1, max_column + 1):
        source = ws.cell(row=source_row, column=column)
        target = ws.cell(row=target_row, column=column)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)
        if source.border:
            target.border = copy.copy(source.border)


def format_summary_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    summary_row: int,
    date_columns: list[int],
    total_column: int,
) -> None:
    first_column = min(date_columns)
    last_column = max(total_column, max(date_columns))
    for column in range(first_column, last_column + 1):
        cell = ws.cell(row=summary_row, column=column)
        if not isinstance(cell, MergedCell):
            cell.fill = copy.copy(GREEN_FILL)


def clear_summary_values(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    summary_row: int,
    date_columns: list[int],
    total_column: int,
) -> None:
    for column in [*date_columns, total_column]:
        cell = ws.cell(row=summary_row, column=column)
        if not isinstance(cell, MergedCell):
            cell.value = None


def collect_day_times(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    column: int,
    first_time_row: int,
    last_time_row: int,
) -> list[int]:
    times: list[int] = []
    for row in range(first_time_row, last_time_row + 1):
        times.extend(extract_time_minutes(ws.cell(row=row, column=column).value))
    return times


def choose_total_column(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    summary_row: int,
    date_columns: list[int],
) -> int:
    last_date_column = max(date_columns)
    rightmost_existing_value = None
    for column in range(last_date_column + 1, ws.max_column + 1):
        value = ws.cell(row=summary_row, column=column).value
        if value not in (None, ""):
            rightmost_existing_value = column

    return rightmost_existing_value or last_date_column + 2


def should_write_zero(existing_value: object) -> bool:
    return existing_value not in (None, "")


def process_employee_section(
    ws: openpyxl.worksheet.worksheet.Worksheet, header_row: int
) -> bool:
    date_row = header_row + 1
    if date_row > ws.max_row:
        return False

    date_columns = find_date_columns(ws, date_row)
    if not date_columns:
        return False

    next_header = find_next_employee_header(ws, header_row)
    search_start = date_row + 1
    search_end = (next_header - 1) if next_header else ws.max_row
    if search_start > search_end:
        return False

    summary_row = find_green_summary_row(ws, search_start, search_end, date_columns)
    created_summary_row = False

    if summary_row is None:
        summary_row = find_first_blank_attendance_row(ws, search_start, search_end, date_columns)

    if summary_row is None:
        if next_header:
            summary_row = next_header
            ws.insert_rows(summary_row)
        else:
            summary_row = search_end + 1
        created_summary_row = True

    total_column = choose_total_column(ws, summary_row, date_columns)

    if created_summary_row and summary_row > 1:
        copy_row_style(ws, summary_row - 1, summary_row, total_column)
        clear_summary_values(ws, summary_row, date_columns, total_column)

    format_summary_row(ws, summary_row, date_columns, total_column)

    first_time_row = date_row + 1
    last_time_row = summary_row - 1
    employee_total = 0.0
    wrote_any_hours = False

    for column in date_columns:
        target_cell = ws.cell(row=summary_row, column=column)
        if isinstance(target_cell, MergedCell):
            continue

        existing_value = target_cell.value
        times = collect_day_times(ws, column, first_time_row, last_time_row)
        hours = calculate_hours_from_times(times)

        if hours is None:
            target_cell.value = 0 if should_write_zero(existing_value) else None
        else:
            rounded_hours = round(hours, 2)
            target_cell.value = rounded_hours
            target_cell.number_format = "0.00"
            employee_total += rounded_hours
            wrote_any_hours = True

    total_cell = ws.cell(row=summary_row, column=total_column)
    if not isinstance(total_cell, MergedCell):
        total_cell.value = round(employee_total, 2) if wrote_any_hours else None
        total_cell.number_format = "0.00"
        total_cell.fill = copy.copy(GREEN_FILL)

    return True


def remove_sheet_protection_and_validations(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    ws.protection.sheet = False
    ws.protection.enable()
    ws.protection.sheet = False
    ws.data_validations = DataValidationList()


def keep_only_original_record(wb: openpyxl.Workbook) -> None:
    for sheet_name in list(wb.sheetnames):
        if sheet_name != TARGET_SHEET:
            del wb[sheet_name]


def calculate_payroll_workbook(uploaded_bytes: bytes) -> tuple[bytes, int]:
    wb = openpyxl.load_workbook(BytesIO(uploaded_bytes))

    if TARGET_SHEET not in wb.sheetnames:
        raise ValueError(f'The workbook must contain a worksheet named exactly "{TARGET_SHEET}".')

    ws = wb[TARGET_SHEET]
    remove_sheet_protection_and_validations(ws)

    processed_sections = 0
    for header_row in reversed(find_employee_header_rows(ws)):
        if process_employee_section(ws, header_row):
            processed_sections += 1

    keep_only_original_record(wb)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), processed_sections


def main() -> None:
    st.set_page_config(page_title="Payroll Calculator", layout="centered")

    st.title("Payroll Calculator")
    st.write("Upload your biometric Excel file and download an updated workbook with calculated payroll hours.")

    uploaded_file = st.file_uploader("Upload your biometric Excel file", type=["xlsx", "xlsm"])

    if uploaded_file is None:
        return

    try:
        with st.spinner("Processing Original record..."):
            updated_file, processed_count = calculate_payroll_workbook(uploaded_file.getvalue())

        if processed_count == 0:
            st.warning("No employee sections with date rows were found on Original record.")
        else:
            st.success(f"Payroll hours calculated successfully for {processed_count} employee section(s).")

        st.download_button(
            "Download updated Original record",
            data=updated_file,
            file_name=OUTPUT_FILENAME,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except ValueError as exc:
        st.error(str(exc))
    except Exception as exc:  # Streamlit should show a friendly message instead of a blank failure.
        st.error("The file could not be processed. Please check that it is a valid Excel workbook.")
        st.exception(exc)


if __name__ == "__main__":
    main()
